"""Arma datos/Ultimo pedido inventario.xlsx desde datos/inventario-corregido.xlsx (19-oct-2026).

    python -I scripts/armar_pedido_inventario.py datos/inventario-corregido.xlsx "datos/Ultimo pedido inventario.xlsx"
"""
import sys, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
src, out = sys.argv[1], sys.argv[2]
amar = PatternFill("solid", fgColor="FFF2CC"); neg = Font(bold=True)
wrap = Alignment(wrap_text=True, vertical="top")

preguntas = [
 ("Pispi Blend de Tintas (W2399)", "La caja es de 6, pero las 5 cajas se habían contado como de 3. ¿Hay 15 o 30 botellas?"),
 ("Estuche Animal Orgánico Colección x6 (W716)", "Hay 1 estuche. ¿Lo cargamos como 1 unidad o como 6 botellas?"),
 ("Salentein Espumante Magnum (W4396)", "En la planilla la caja es de 4 (1 caja + 5 sueltas = 9) y en el sistema, de 6. ¿Cuál es?"),
 ("Doña Paula Selección de Bodega Malbec (W1067)", "Aparece en dos filas, con 18 y 12. ¿Se suman (30) o vale una sola?"),
 ("Trivento Golden Reserva Cabernet Sauvignon (45 botellas)", "Confirmaste W3313, pero en el sistema W3313 es el Trivento Reserve común. El Golden CS es W4412. ¿Lo cargamos en W4412?"),
 ("Trivento Golden Reserva Cabernet Franc (18 botellas)", "El Golden CF ya existe en el sistema: es W3366. ¿Lo cargamos ahí y dejamos W3303 (La Salonera) como está?"),
 ("Owner Malbec Red Blend (24 botellas)", "Le pusiste W4681, pero ese código es el Catena Malbec Argentino. ¿Qué código le damos al Owner?"),
 ("Catena Malbec Argentino (40 botellas)", "Pusiste W4686 y que el W654 es el que viene por 4. ¿Lo cargamos en W4681, que es el de caja de 6?"),
 ("Raquis Las Bases Malbec 2021 (1 caja)", "¿Es el W2919? En el sistema viene en caja de 3: ¿son 3 o 6 botellas?"),
 ("Birth of Catena (1 caja + 4 sueltas)", "En el sistema hay dos iguales: W44505 (caja de 4) y W44506 (caja de 6). ¿Cuál es?"),
 ("Casalta Clásico vidrio x400 (13 unidades)", "¿Qué código tiene? El W4186 es el Casalta Reducción."),
 ("Chachingo Porter (1 unidad)", "¿Qué código tiene?"),
 ("Chachingo Amber (1 unidad)", "¿Qué código tiene?"),
 ("Casarena DNA Cabernet Sauvignon 2017 (1 caja)", "¿Es el W542? Ya tiene 3 botellas cargadas de otra fila: ¿se suman?"),
 ("Sylvestra Pinot (2 sueltas) → W453", "Ya tiene 5 cargadas de otra fila. ¿Se suman (7)?"),
 ("Las Perdices Syrah Viognier (5 sueltas) → W1999", "W1999 es el Las Perdices Viognier y ya tiene 3. ¿Es el mismo vino y se suman (8)?"),
 ("Codorníu Tradicional Semi Sec (2 sueltas) → W2972", "W2972 es el María Codorníu Método Tradicional y ya tiene 20. ¿Es el mismo y se suman (22)?"),
 ("Codorníu Clásico (5 sueltas) → W2974", "W2974 es el María Codorníu Brut Nature y ya tiene 8. ¿Es el mismo y se suman (13)?"),
]

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Preguntas"
ws.append(["Contestá en la columna amarilla. Son las respuestas que nos faltan para terminar de cargar el inventario."])
ws["A1"].font = Font(italic=True); ws.append([])
ws.append(["#", "Producto", "Pregunta", "RESPUESTA"])
for c in ws[3]: c.font = neg
for i, (p, q) in enumerate(preguntas, 1):
    ws.append([i, p, q, None]); ws.cell(ws.max_row, 4).fill = amar
    for c in ws[ws.max_row]: c.alignment = wrap
for col, w in zip("ABCD", [4, 48, 70, 30]): ws.column_dimensions[col].width = w
ws.freeze_panes = "A4"

# Productos nuevos: codigo que dio Federico + nombre + caja de su planilla si la tenia
inv = openpyxl.load_workbook(src, data_only=True)
nuevos = []
ARREGLOS = {"SELECCI�N":"SELECCIÓN","OTO�O":"OTOÑO","Imag�nate":"Imagínate","VI�A":"VIÑA","ESP�MANTE":"ESPUMANTE","DO�A":"DOÑA","AQU�":"AQUÍ"}
def limpio(s):
    s = " ".join(str(s or "").split())
    for a, b in ARREGLOS.items(): s = s.replace(a, b)
    return s
def filas(hoja):
    return [r for r in inv[hoja].iter_rows(min_row=4, values_only=True) if r[0] is not None or r[1]]
for r in filas("Codigo de otro producto"):
    if r[8] in ("W4762","W4763","W4764","W4765","W4766","W4767"): nuevos.append((r[8], r[2], r[3], r[4], r[5], r[6]))
for r in filas("No esta en el sistema"):
    cod = r[8]
    if cod in ("W4377","W4166","W4768","W4769","W4770","W4185","W429","W4745","W4746","W4747","W4748","W4750","W4757","W4758","W4164"):
        nuevos.append((cod, r[2], r[3], r[4], r[5], r[6]))
for r in filas("Sin codigo"):
    if r[8] and r[8][1:].isdigit() and 4771 <= int(r[8][1:]) <= 4816: nuevos.append((r[8], r[2], r[3], None, r[5], r[6]))

ws2 = wb.create_sheet("Productos nuevos")
ws2.append(["Los vamos a dar de alta con el código que nos pasaste y precio «a consultar»: el precio se lo ponen ustedes después desde Productos → Editar. Necesitamos cuántas botellas trae la caja. Si la columna ya tiene número, es lo que decía tu planilla: corregilo solo si está mal."])
ws2["A1"].font = Font(italic=True); ws2.merge_cells("A1:H1"); ws2["A1"].alignment = wrap; ws2.row_dimensions[1].height = 45
ws2.append([])
ws2.append(["Código", "Bodega", "Producto", "Cajas contadas", "Sueltas contadas", "BOTELLAS POR CAJA", "¿SE VENDE SUELTO? (sí/no)"])
for c in ws2[3]: c.font = neg
for cod, bod, prod, xc, cj, su in sorted(nuevos, key=lambda n: n[0]):
    ws2.append([cod, limpio(bod), limpio(prod), cj or 0, su or 0, xc, None])
    ws2.cell(ws2.max_row, 6).fill = amar; ws2.cell(ws2.max_row, 7).fill = amar
for col, w in zip("ABCDEFG", [9, 18, 55, 9, 9, 12, 14]): ws2.column_dimensions[col].width = w
ws2.freeze_panes = "A4"
wb.save(out)
print(len(preguntas), "preguntas,", len(nuevos), "productos nuevos,", sum(1 for n in nuevos if n[3]), "con caja de la planilla")
